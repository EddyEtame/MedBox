"""Build the crew's defence deck (PPTX): five minutes, straight to the point.

Ten slides, dark, one idea per slide, in our own voice: what we had to do,
what we did, then MedBox (1A, Eddy and Brad), ARIA / PsychoSpace (1B,
Davidson, Anthony, Frederic, Merove), what comes next, one sentence.
Document tooling only (python-pptx); output goes to .build\\deliverables,
ignored by Git. Page captures from .build\\deliverables\\shots are used
when present.

    .venv\\Scripts\\python tools\\build_deck.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / ".build" / "deliverables"
SHOTS = OUT / "shots"

BG = RGBColor(0x05, 0x0F, 0x17)
PANEL = RGBColor(0x0B, 0x1B, 0x25)
INK = RGBColor(0xEA, 0xFB, 0xF8)
MUTED = RGBColor(0x8F, 0xA8, 0xB2)
ACCENT = RGBColor(0x7B, 0xE8, 0xD3)
WARN = RGBColor(0xF0, 0xC4, 0x59)
CRIT = RGBColor(0xFF, 0x5C, 0x6E)
PLUM = RGBColor(0xC9, 0xA7, 0xF0)

W, H = Inches(13.333), Inches(7.5)
TOTAL = 10


def count_tests() -> int:
    """The number pytest reports on the last full run (the log in
    .build/logs), which counts parametrised cases; the source count is the
    fallback when no log exists."""
    logs = sorted((ROOT / ".build" / "logs").glob("pytest-*.log"), key=lambda f: f.stat().st_mtime)
    for log in reversed(logs):
        m = re.findall(r"(\d+) passed", log.read_text(encoding="utf-8", errors="replace"))
        if m:
            return int(m[-1])
    n = 0
    for f in (ROOT / "tests").glob("test_*.py"):
        n += len(re.findall(r"^\s*def test_", f.read_text(encoding="utf-8", errors="replace"), re.M))
    return n


def count_scenarios() -> int:
    return len(list((ROOT / "scenarios").glob("*.yaml")))


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def slide(self, kicker: str, title: str, accent=ACCENT):
        self.n += 1
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = BG
        self.text(s, kicker.upper(), Inches(0.6), Inches(0.35), Inches(11), Inches(0.4), size=12, color=accent, bold=True, spacing=2)
        self.text(s, title, Inches(0.6), Inches(0.7), Inches(12), Inches(1.0), size=30, color=INK, bold=True)
        self.text(s, f"ÉQUIPAGE HORIZON 2080  ·  1A MEDBOX  ·  1B ARIA          {self.n} / {TOTAL}", Inches(0.6), Inches(6.95), Inches(12), Inches(0.35), size=10, color=MUTED, spacing=2)
        return s

    @staticmethod
    def text(s, txt, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=0, font="Segoe UI"):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(txt.split("\n")):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = align
            run = par.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = font
            run.font.color.rgb = color
            if spacing:
                run.font._element.set("spc", str(spacing * 100))
        return tb

    @staticmethod
    def panel(s, x, y, w, h, fill=PANEL, line=None):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shp.adjustments[0] = 0.06
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = Pt(1)
        shp.shadow.inherit = False
        return shp

    def card(self, s, x, y, w, h, head, body, head_color=ACCENT, size=13):
        self.panel(s, x, y, w, h)
        self.text(s, head, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), Inches(0.5), size=16, color=head_color, bold=True)
        self.text(s, body, x + Inches(0.25), y + Inches(0.75), w - Inches(0.5), h - Inches(0.9), size=size, color=INK)

    def rows(self, s, items, y0=1.75, step=1.0, head_color=ACCENT, head_w=3.0, size=13):
        for i, (head, body) in enumerate(items):
            y = Inches(y0 + i * step)
            self.panel(s, Inches(0.6), y, Inches(12.1), Inches(step - 0.1))
            self.text(s, head, Inches(0.85), y + Inches(0.13), Inches(head_w), Inches(0.6), size=15, color=head_color, bold=True)
            self.text(s, body, Inches(0.85 + head_w + 0.2), y + Inches(0.14), Inches(12.1 - head_w - 0.7), Inches(step - 0.2), size=size, color=INK)

    def picture(self, s, name, x, y, w):
        f = SHOTS / name
        if f.exists():
            s.shapes.add_picture(str(f), x, y, width=w)
            return True
        return False

    def number(self, s, x, y, w, value, label, color=ACCENT):
        self.text(s, value, x, y, w, Inches(0.9), size=40, color=color, bold=True, align=PP_ALIGN.CENTER)
        self.text(s, label, x, y + Inches(0.9), w, Inches(0.6), size=12, color=MUTED, align=PP_ALIGN.CENTER)

    def save(self, path: Path):
        self.prs.save(str(path))


def build() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    tests, scenarios = count_tests(), count_scenarios()
    d = Deck()

    # 1 — title
    s = d.slide("Workshop 2026 · B3 · Horizon 2080 · HumanTech & HealthTech spatiales", "Deux solutions pour le vaisseau-monde")
    d.text(s, "1A · MedBox — le référent médical du bord, pour le corps\n1B · ARIA / PsychoSpace — l’assistant de bord, pour l’esprit",
           Inches(0.6), Inches(1.8), Inches(12), Inches(1.3), size=20, color=ACCENT)
    d.text(s, "Six personnes. Deux prototypes qui tournent. Cinq minutes.",
           Inches(0.6), Inches(3.1), Inches(11.5), Inches(0.9), size=16, color=INK)
    if not d.picture(s, "ship.png", Inches(0.6), Inches(4.05), Inches(4.9)):
        d.panel(s, Inches(0.6), Inches(4.05), Inches(4.9), Inches(2.6))
    d.text(s, "MedBox : Eddy · Brad\nARIA / PsychoSpace : Davidson · Anthony · Frederic · Merove", Inches(6.5), Inches(5.6), Inches(6.3), Inches(0.9), size=14, color=MUTED, align=PP_ALIGN.RIGHT)

    # 2 — what we had to do
    s = d.slide("La commande", "Ce que nous avions à faire")
    d.rows(s, [("Deux projets", "Un équipage de six doit livrer deux solutions justifiées, avec un effet réel sur la vie à bord, et des pistes pour les années suivantes."),
               ("Un vaisseau sans la Terre", "Des décennies de voyage, une liaison inutilisable : tout doit fonctionner à bord, sans réseau, et rester compréhensible par l’équipage."),
               ("Le pilier HumanTech & HealthTech", "Mesurer, apprendre la ligne de base de chacun, suivre l’historique, comparer, analyser avec une IA locale, agir. Pour le corps et pour l’esprit."),
               ("Les livrables", "Un prototype qui tourne, le code sur Git, un dossier technique, ce support. Et une démonstration qui tient en cinq minutes.")],
           y0=1.8, step=1.15, head_w=3.3, size=14)

    # 3 — what we did
    s = d.slide("Le résultat", "Ce que nous avons fait")
    d.card(s, Inches(0.6), Inches(1.85), Inches(5.95), Inches(3.35), "1A · MEDBOX — EDDY, BRAD",
           "Une station médicale qui mesure quarante personnes dix fois par seconde, décide qui isoler, l’explique et le dit à voix haute, dans la langue de chacun.\n\nLivrée en un dossier de 1,6 Go qui se lance en quatre secondes, sans réseau.", size=14)
    d.card(s, Inches(6.75), Inches(1.85), Inches(5.95), Inches(3.35), "1B · ARIA / PSYCHOSPACE — DAVIDSON, ANTHONY, FREDERIC, MEROVE",
           "Un assistant de bord qui suit le bien-être de chaque astronaute, repère une dérive avant qu’elle ne devienne une crise, consulte les procédures du vaisseau et propose une action.\n\nLocal, explicable, l’humain garde la main.", PLUM, size=14)
    d.panel(s, Inches(0.6), Inches(5.4), Inches(12.1), Inches(1.0))
    d.text(s, "EN COMMUN", Inches(0.85), Inches(5.5), Inches(2.4), Inches(0.4), size=12, color=WARN, bold=True, spacing=1)
    d.text(s, "Une ligne de base par personne · des règles explicites avant tout modèle · une IA locale qui formule et ne décide pas · zéro octet vers le réseau",
           Inches(3.3), Inches(5.5), Inches(9.2), Inches(0.85), size=13, color=INK)

    # 4 — MedBox: what we built
    s = d.slide("1A · MedBox", "Nous avons construit le référent médical du bord")
    d.card(s, Inches(0.6), Inches(1.9), Inches(3.9), Inches(3.9), "IL MESURE",
           "Quarante membres, cinq constantes, dix fois par seconde, comparées à la plage habituelle de chacun. Une seule valeur ne suffit jamais.")
    d.card(s, Inches(4.7), Inches(1.9), Inches(3.9), Inches(3.9), "IL DÉCIDE",
           "Un score NEWS2 par des règles que n’importe qui peut relire. Il nomme ce que les mesures montrent, décide l’isolement, attribue une zone et ses couchettes.", WARN)
    d.card(s, Inches(8.8), Inches(1.9), Inches(3.9), Inches(3.9), "IL PARLE",
           "À la personne et à l’équipage, à l’écran et à voix haute, en français ou en anglais. Il écoute jusqu’au bout de la phrase, s’éveille sur son nom, répond, fait la ronde du vaisseau.", CRIT)
    d.text(s, "Quatre pages, sept serveurs, un espace personnel par membre, une vue 3D du vaisseau pièce par pièce.", Inches(0.6), Inches(6.1), Inches(12), Inches(0.6), size=13, color=MUTED)

    # 5 — MedBox: demo
    s = d.slide("1A · MedBox", "La démonstration, en cinq temps")
    d.rows(s, [("1 · Lancement", "Double-clic : station, six espaces et modèle prêts en quatre secondes. Le référent se présente ; consentement à la voix ; langue."),
               ("2 · Équipage", "La semaine de l’équipe, le membre en forme et ses habitudes, les activités du jour."),
               ("3 · Contamination", "Six membres se dégradent ; messages sonores ; « Lire » : carte, vaisseau 3D, la pièce qui s’illumine, la décision lue."),
               ("4 · Espace personnel", "Le référent reconnaît la personne, lit son évaluation à la deuxième personne, répond à sa question parlée."),
               ("5 · Panne", "Nous arrêtons le modèle : la surveillance continue, l’écran le dit. Nous le relançons : le référent revient.")],
           y0=1.75, step=1.0)

    # 6 — MedBox: what it cannot do
    s = d.slide("1A · MedBox", "Ce que le référent ne peut pas faire, par construction")
    d.text(s, "Deux chemins qui ne se mélangent jamais : les mesures, le score et l’isolement d’un côté, sans modèle ; le langage de l’autre, sous schéma et validateur.",
           Inches(0.6), Inches(1.7), Inches(12), Inches(0.8), size=14, color=INK)
    d.rows(s, [("Décider du score", "Le score et la proposition d’isolement viennent de règles explicites. Le modèle formule, il ne calcule rien."),
               ("Prescrire", "Aucun médicament, aucune dose ne passe le validateur. Jamais."),
               ("Inventer", "Une condition n’est nommée que si les mesures la montrent ; un isolé n’existe que dans les registres de la station."),
               ("Décider seul", "Confirmation et levée d’isolement sont humaines. Le consentement micro est parlé, révocable, et rien de l’audio n’est conservé."),
               ("Tomber en silence", "Si le modèle s’arrête, la surveillance, les décisions et les pages continuent, et l’écran le dit.")],
           y0=2.55, step=0.84, head_color=CRIT, head_w=3.0)

    # 7 — MedBox: numbers
    s = d.slide("1A · MedBox", "Vérifié sur la machine de soutenance")
    cols = [("40", "membres suivis"), (str(scenarios), "scénarios rejouables"), (str(tests), "tests verts"),
            ("4 s", "lancement à froid"), ("7", "serveurs, un dossier"), ("0", "octet vers le réseau")]
    for i, (v, l) in enumerate(cols):
        x = Inches(0.6 + (i % 3) * 4.1)
        y = Inches(1.9 + (i // 3) * 2.2)
        d.panel(s, x, y, Inches(3.9), Inches(2.0))
        d.number(s, x, y + Inches(0.25), Inches(3.9), v, l, color=ACCENT if i != 5 else WARN)

    # 8 — ARIA: what they built
    s = d.slide("1B · ARIA / PsychoSpace", "Ils ont construit l’assistant qui voit la dérive avant la crise", PLUM)
    d.card(s, Inches(0.6), Inches(1.9), Inches(3.9), Inches(3.9), "IL OBSERVE",
           "Check-ins quotidiens : sommeil, humeur, fatigue, stress, isolement ; à terme des capteurs ESP32. Un historique local par équipier.", PLUM)
    d.card(s, Inches(4.7), Inches(1.9), Inches(3.9), Inches(3.9), "IL COMPARE",
           "Une baseline personnelle, des écarts persistants, des signaux convergents. Des règles explicites, pas une décision opaque du modèle.", WARN)
    d.card(s, Inches(8.8), Inches(1.9), Inches(3.9), Inches(3.9), "IL ACCOMPAGNE",
           "Il explique l’observation, pose des questions, consulte les procédures du vaisseau (RAG local) et propose une action avec un niveau de priorité. Jamais de diagnostic.", CRIT)
    d.text(s, "Python / FastAPI · SQLite · Ollama (Llama 3.2) et embeddings locaux · interface web · ESP32 prévu.", Inches(0.6), Inches(6.1), Inches(12), Inches(0.6), size=13, color=MUTED)

    # 9 — ARIA: scenario and state
    s = d.slide("1B · ARIA / PsychoSpace", "Le scénario central, et ce qui tourne", PLUM)
    d.card(s, Inches(0.6), Inches(1.85), Inches(5.95), Inches(3.5), "LA DÉRIVE PROGRESSIVE",
           "Sommeil proche de 7 h 30, humeur haute, fatigue faible. Puis, jour après jour : le sommeil baisse, la fatigue monte, l’activité diminue, un retrait social apparaît.\n\n« Ça va, je suis juste fatigué. » ARIA ne contredit pas : il compare à l’historique, explique la dérive, demande si un événement l’explique, propose une action cohérente.", PLUM, size=12)
    d.card(s, Inches(6.75), Inches(1.85), Inches(5.95), Inches(3.5), "ÉTAT DU PROTOTYPE",
           "IA locale : fonctionnelle · Conversation : fonctionnelle · Suivi bien-être : fonctionnel · Aide médicale : en démonstration · RAG local : V2 · Vue équipage : prototype · Capteurs : simulés · Moteur de dérive avancé : architecture cible.\n\nRésultat clé : l’assistance reste locale, garde son contexte et exploite une base documentaire embarquée sans API cloud.", WARN, size=12)
    d.panel(s, Inches(0.6), Inches(5.55), Inches(12.1), Inches(0.9))
    d.text(s, "Réponse structurée : PRIORITÉ · OBSERVATION · CONTEXTE · ACTION PROPOSÉE · SUIVI · SOURCE        Dépôt : github.com/ANTHONYSITCH/Psychospace",
           Inches(0.85), Inches(5.7), Inches(11.7), Inches(0.7), size=12, color=MUTED)

    # 10 — next, and one sentence
    s = d.slide("La suite", "Ce qui vient, et ce que nous retenons")
    d.rows(s, [("MedBox", "Coque du vaisseau en 3D, ronde horaire parlée, brief du matin, tête de mesure ESP32, mode basse consommation."),
               ("ARIA", "Capteurs ESP32 réels, RAG enrichi et versions des procédures, mémoire personnalisée, déploiement résilient."),
               ("Ensemble", "Un isolement décidé par MedBox devient un suivi de bien-être pour ARIA ; une dérive vue par ARIA appelle une mesure de MedBox.")],
           y0=1.75, step=0.95, head_color=WARN, head_w=2.4)
    d.panel(s, Inches(0.6), Inches(4.75), Inches(12.1), Inches(1.9), fill=PANEL, line=ACCENT)
    d.text(s, "« Un vaisseau-monde n’a pas de médecin ni de psychologue à bord. Il a MedBox et ARIA : deux systèmes locaux qui mesurent, "
              "expliquent et accompagnent, et qui continuent quand la Terre ne répond plus. »",
           Inches(1.0), Inches(4.95), Inches(11.3), Inches(1.6), size=17, color=INK)

    path = OUT / "Workshop2026-B3-Pres-Equipage.pptx"
    d.save(path)
    return path


def main() -> int:
    print(build())
    return 0


if __name__ == "__main__":
    sys.exit(main())
